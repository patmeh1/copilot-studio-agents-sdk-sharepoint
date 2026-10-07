"""Build the Agents SDK tutorial, use cases / test scenarios (markdown), the Excel test plan and the JSON catalog.

    python tools/build_sdk_docs.py
Content comes from tools/sdk_content.py. Recalculate the workbook afterwards (formulas):
    python <xlsx skill>/scripts/recalc.py docs/Agents-SDK-SharePoint-Test-Plan.xlsx
"""
import json
import os
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import sdk_content as c

DOCS = os.path.join(c.ROOT, "docs")
TUTORIAL_MD = "Agents-SDK-SharePoint-Tutorial.md"
SCENARIOS_MD = "Agents-SDK-SharePoint-Use-Cases-and-Test-Scenarios.md"
HTML = "Agents-SDK-SharePoint-Tutorial-and-Test-Scenarios.html"
XLSX = "Agents-SDK-SharePoint-Test-Plan.xlsx"
CATALOG = os.path.join("test-results", "agents-sdk-test-catalog.json")
ICON = {"pass": "✅", "limit": "⛔", "finding": "⚠️", "notrun": "⬜"}
UC_ICON = {"validated": "✅", "design": "📘", "validate": "🔎"}


def slug(text):
    return re.sub(r"\s", "-", re.sub(r"[^\w\s-]", "", text.lower()).strip())


def plain(text):
    """Strip inline markup for Excel; literal asterisks inside `code` are kept."""
    codes = []

    def keep(m):
        codes.append(m.group(1))
        return f"\x00{len(codes) - 1}\x00"

    text = re.sub(r"`([^`]+)`", keep, text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "")
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", text)
    return re.sub(r"\x00(\d+)\x00", lambda m: codes[int(m.group(1))], text)


def tc_link(tid, same_page=False):
    t = next(x for x in c.TESTS if x["id"] == tid)
    return f"[{tid}]({'' if same_page else SCENARIOS_MD}#{slug(t['id'] + ' ' + t['title'])})"


def companion(current):
    links = [("Tutorial", TUTORIAL_MD), ("Use cases and test scenarios", SCENARIOS_MD), ("Test plan (Excel)", XLSX),
             ("Web page", c.SITE_URL + HTML), ("Full guide and findings", "Agents-SDK-SharePoint-Guide.md")]
    return " · ".join(f"[{n}]({p})" for n, p in links if p != current)


def tutorial_md():
    out = ["# Tutorial: an authenticated Copilot Studio agent on any SharePoint site with the Microsoft 365 Agents SDK", "",
           f"> Every step was executed and verified on **{c.TESTED_ON}** in a Microsoft demo tenant ({c.STACK}, web part package "
           f"{c.PACKAGE_VERSION}).", f"> {companion(TUTORIAL_MD)}", "",
           "## What you will build", "",
           "One SharePoint Framework (SPFx) web part, deployed once for the whole tenant, that shows an authenticated Copilot Studio "
           "agent inline on any page of any site. Users are signed in with their existing Microsoft Entra ID session (single "
           "sign-on), the agent knows who they are, and Copilot Studio sharing decides who may chat.", "",
           "```mermaid", "flowchart LR",
           "    U[User signed in to SharePoint] --> P[Page on any site]",
           "    P --> W[Copilot Studio Agent web part]",
           "    W -- MSAL single sign-on --> E[Microsoft Entra ID]",
           "    E -- token for CopilotStudio.Copilots.Invoke --> W",
           "    W -- Agents SDK + connection string --> C[Copilot Studio agent]",
           "    C -- sharing check --> R[Answers in the page]", "```", "",
           "## Before you start", "",
           "| Need | Details |", "|---|---|",
           "| Copilot Studio | Maker access to the environment; Copilot Studio capacity (Copilot Credits) - users do not need a Microsoft 365 Copilot license |",
           "| Microsoft Entra ID | Permission to register an app and grant tenant-wide consent (Cloud Application Administrator or higher) |",
           "| SharePoint | SharePoint Administrator and site admin of the tenant app catalog |",
           "| Workstation | Windows, PowerShell 7, Azure CLI, Node.js 22 (SPFx 1.21.1 needs >= 22.14 < 23) and a clone of this repository |",
           "| Time | About 60-90 minutes the first time. Adding the agent to another site takes about 2 minutes. |", ""]
    n = 0
    for pi, part in enumerate(c.TUTORIAL, 1):
        out += [f"## Part {pi} - {part['part']}", ""]
        for st in part["steps"]:
            n += 1
            out += [f"### Step {n} - {st['title']}", "", f"*Why:* {st['why']}", ""]
            out += [f"{i}. {d}" for i, d in enumerate(st["do"], 1)] + [""]
            if st.get("code"):
                lang, code = st["code"]
                out += [f"```{lang}", code, "```", ""]
            out += [f"> **Checkpoint:** {st['check']}", ""]
            for img, cap in st["images"]:
                out += [f"![{cap}](images/{img})", f"*{cap}*", ""]
    out += ["## Troubleshooting", "",
            "| What you see | What to do |", "|---|---|",
            "| *Sign in to chat with this agent* | Silent sign-on was blocked (third-party cookies, Conditional Access, embedding host). Select **Sign in with Microsoft**. |",
            "| *The sign-in popup was blocked...* | Allow popups for `https://<tenant>.sharepoint.com` and select **Try again**. Package 1.0.2.0+ opens the popup on the first click, also with several agents on a page. |",
            "| *Redirect URI mismatch* (`AADSTS50011`) | Add `https://<tenant>.sharepoint.com` (no trailing slash) as a **Single-page application** redirect URI of the app. |",
            "| *Consent is missing* (`AADSTS65001`) | Grant admin consent for **Power Platform API > CopilotStudio.Copilots.Invoke**. |",
            "| *You don't have access to talk to this bot* | Share the agent with the user or with a security group the user belongs to. |",
            "| *This action doesn't support agents built with the GitHub Copilot harness* | Use an agent built on the standard harness. |",
            "| App catalog: *you need to add yourself as a site admin* | SharePoint admin center > Active sites > Apps > Membership > Site admins. |", "",
            "## Next", "",
            f"* Run the full set of [test scenarios]({SCENARIOS_MD}#test-scenarios) and record your results in the [test plan]({XLSX}).",
            "* Background, findings and the comparison with the SharePoint channel + Agent link pattern: [guide](Agents-SDK-SharePoint-Guide.md).", ""]
    return "\n".join(out)


def scenarios_md():
    k = c.counts()
    out = ["# Microsoft 365 Agents SDK in SharePoint - use cases and test scenarios", "",
           f"> Tested on **{c.TESTED_ON}** in a Microsoft demo tenant with 4 test users, 3 sites, 2 agents and the maker as control "
           f"(web part {c.PACKAGE_VERSION}). **{len(c.TESTS)} test scenarios:** {k['pass']} passed, {k['limit']} confirmed limitation, "
           f"{k['finding']} security finding, {k['notrun']} to validate in your tenant.", f"> {companion(SCENARIOS_MD)}", "",
           "## Use cases", "",
           "| ID | Use case | Status | Test scenarios |", "|---|---|---|---|"]
    for uc in c.USE_CASES:
        tests = ", ".join(tc_link(t, same_page=True) for t in uc["tests"]) or "-"
        out.append(f"| [{uc['id']}](#{slug(uc['id'] + ' ' + uc['title'])}) | {uc['title']} | {UC_ICON[uc['status']]} {c.UC_STATUS[uc['status']][0]} | {tests} |")
    out.append("")
    for uc in c.USE_CASES:
        out += [f"### {uc['id']} {uc['title']}", "",
                f"* **Scenario:** {uc['scenario']}", f"* **How the Agents SDK web part does it:** {uc['how']}",
                f"* **Evidence:** {uc['evidence']}", f"* **Status:** {c.UC_STATUS[uc['status']][0]}", ""]
    out += ["### Not supported", ""] + [f"* **{a}** - {b}" for a, b in c.NOT_SUPPORTED] + [""]
    out += ["## Test setup", "",
            "Agent A is a standard-harness agent shared with User A and User B. Agent B (Weather) is a GitHub Copilot harness agent "
            "shared with User A and User C. User D has no agent. Every site page hosts Agent A and Agent B through the Agents SDK "
            "web part plus Agent B's hosted embed for comparison. Test users have Microsoft 365 E5 (no Teams), no Copilot licenses, "
            "and MFA.", "",
            "| Persona | Agent sharing | Site A | Site B | Site C |", "|---|---|---|---|---|"]
    out += [f"| {' | '.join(p)} |" for p in c.PERSONAS] + [""]
    out += ["## Test scenarios", "",
            f"Run the **core set** ({', '.join(c.CORE_SET)}) after every change. Priority P1 = must pass before go-live, "
            "P2 = should pass, P3 = optional. *Automation* names the script in `tests/agents-sdk-matrix`.", "",
            "| ID | Area | Scenario | Priority | Result in the demo tenant |", "|---|---|---|---|---|"]
    for t in c.TESTS:
        out.append(f"| {tc_link(t['id'], same_page=True)} | {t['area']} | {t['title']} | {t['priority']} | "
                   f"{ICON[t['status']]} {c.STATUS[t['status']][0]} |")
    out.append("")
    for area in c.AREAS:
        out += [f"### {area}", ""]
        for t in [x for x in c.TESTS if x["area"] == area]:
            ev = " · ".join(f"[{os.path.basename(e)}]({e})" for e in t["evidence"]) or "-"
            steps = "<br>".join(f"{i}. {s}" for i, s in enumerate(t["steps"], 1))
            out += [f"#### {t['id']} {t['title']}", "", "| | |", "|---|---|",
                    f"| **Priority** | {t['priority']} |", f"| **Preconditions** | {t['pre']} |", f"| **Steps** | {steps} |",
                    f"| **Expected result** | {t['expected']} |", f"| **Result in the demo tenant** | {t['actual']} |",
                    f"| **Status** | {ICON[t['status']]} {c.STATUS[t['status']][0]} |", f"| **Automation** | {t['auto']} |",
                    f"| **Evidence** | {ev} |", ""]
    return "\n".join(out)


THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", start_color="2B2A29")
SOFT_FILL = PatternFill("solid", start_color="F2F2F2")
FILLS = {"Pass": "E3F4E8", "Confirmed limitation": "FBE5E5", "Security finding": "FDF0D8", "Fail": "FBE5E5", "Blocked": "FDF0D8"}


def _font(**kw):
    return Font(name="Arial", size=kw.pop("size", 10), **kw)


def _table(ws, row, headers, rows, widths):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    for r, values in enumerate(rows, row + 1):
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = _font()
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    return row + len(rows)


def _fit_rows(ws, first_row, last_row, header_row=None, line_pt=12.75):
    """Explicit row heights from the wrapped text, so nothing is clipped (openpyxl does not auto-fit)."""
    widths = {col: dim.width for col, dim in ws.column_dimensions.items()}
    for r in range(first_row, last_row + 1):
        lines = 1
        for cell in ws[r]:
            if cell.value is None or isinstance(cell.value, (int, float)) or str(cell.value).startswith("="):
                continue
            per_line = max(1.0, (widths.get(cell.column_letter, 10) or 10) * 1.05)
            n = sum(max(1, -(-len(part) // int(per_line))) for part in str(cell.value).split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = lines * line_pt + 5
    if header_row:
        ws.row_dimensions[header_row].height = 30


def _print_setup(ws, title_rows=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows


def workbook(path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    tests = wb.create_sheet("Test scenarios")
    usecases = wb.create_sheet("Use cases")
    setup = wb.create_sheet("Test setup")

    headers = ["ID", "Area", "Priority", "Scenario", "Preconditions", "Steps", "Expected result",
               f"Result in demo tenant ({c.TESTED_ON})", "Demo status", "Automation", "Evidence (repository docs/)",
               "Your result", "Tester", "Date", "Notes"]
    rows = [[t["id"], t["area"], t["priority"], t["title"], plain(t["pre"]),
             "\n".join(f"{i}. {plain(s)}" for i, s in enumerate(t["steps"], 1)), plain(t["expected"]), plain(t["actual"]),
             c.STATUS[t["status"]][0].split(" - ")[0], t["auto"], "\n".join(t["evidence"]), None, None, None, None] for t in c.TESTS]
    last = _table(tests, 1, headers, rows, [8, 18, 8, 34, 38, 44, 40, 48, 16, 20, 34, 12, 14, 12, 30])
    tests.freeze_panes = "E2"
    tests.auto_filter.ref = f"A1:O{last}"
    for col in ("L", "M", "N", "O"):
        for r in range(2, last + 1):
            tests[f"{col}{r}"].fill = PatternFill("solid", start_color="FFF8E1")
    dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not run,N/A"', allow_blank=True)
    dv.prompt, dv.promptTitle = "Pass, Fail, Blocked, Not run or N/A", "Your result"
    tests.add_data_validation(dv)
    dv.add(f"L2:L{last}")
    for col in ("I", "L"):
        for status, color in FILLS.items():
            tests.conditional_formatting.add(f"{col}2:{col}{last}", FormulaRule(formula=[f'{col}2="{status}"'],
                                                                              fill=PatternFill("solid", start_color=color)))

    uc_rows = [[u["id"], u["title"], c.UC_STATUS[u["status"]][0], plain(u["scenario"]), plain(u["how"]), plain(u["evidence"]),
                ", ".join(u["tests"])] for u in c.USE_CASES]
    uc_rows += [["-", plain(a), "Not supported", plain(b), "", "", ""] for a, b in c.NOT_SUPPORTED]
    last_uc = _table(usecases, 1, ["ID", "Use case", "Status", "Scenario", "How the Agents SDK web part does it", "Evidence", "Test scenarios"],
                     uc_rows, [8, 34, 18, 48, 52, 52, 20])
    usecases.freeze_panes = "C2"
    usecases.auto_filter.ref = f"A1:G{last_uc}"

    setup["A1"] = "Test setup used in the demo tenant"
    setup["A1"].font = _font(bold=True, size=12)
    r = _table(setup, 3, ["Persona", "Agent sharing", "Site A", "Site B", "Site C"], [list(p) for p in c.PERSONAS], [20, 22, 14, 14, 14])
    notes = ["Agent A: standard harness, Authenticate with Microsoft, shared with User A and User B.",
             "Agent B (Weather): GitHub Copilot harness, Authenticate with Microsoft, shared with User A and User C.",
             "Each site page hosts Agent A (Agents SDK), Agent B (Agents SDK) and Agent B (hosted embed, comparison).",
             "Test users: Microsoft 365 E5 (no Teams), no Microsoft 365 Copilot or Copilot Studio license, MFA with an authenticator app.",
             f"Web part package {c.PACKAGE_VERSION}; {plain(c.STACK)}.",
             f"Repository: {c.REPO_URL}"]
    for i, n in enumerate(notes, r + 2):
        setup.cell(row=i, column=1, value=n).font = _font()

    s = summary
    s.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        s.column_dimensions[col].width = 15
    s["A1"] = "Microsoft 365 Agents SDK in SharePoint - test plan"
    s["A1"].font = _font(bold=True, size=14)
    s["A2"] = (f"Tested in a Microsoft demo tenant on {c.TESTED_ON}. Record your own results in the yellow columns of "
               "'Test scenarios' (Your result, Tester, Date, Notes); this sheet updates automatically.")
    s["A2"].font = _font(color="5C5C5C")
    s["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A2:G2")
    s.row_dimensions[2].height = 30
    rng_area, rng_demo, rng_mine, rng_prio = "'Test scenarios'!$B$2:$B$200", "'Test scenarios'!$I$2:$I$200", "'Test scenarios'!$L$2:$L$200", "'Test scenarios'!$C$2:$C$200"
    head = ["Area", "Scenarios", "Pass (demo)", "Limitation / finding", "Not run (demo)", "Your pass", "Your fail"]
    _table(s, 4, head, [], [30, 15, 15, 15, 15, 15, 15])
    row = 5
    for area in c.AREAS:
        s.cell(row=row, column=1, value=area)
        s.cell(row=row, column=2, value=f'=COUNTIF({rng_area},A{row})')
        s.cell(row=row, column=3, value=f'=COUNTIFS({rng_area},A{row},{rng_demo},"Pass")')
        s.cell(row=row, column=4, value=f'=COUNTIFS({rng_area},A{row},{rng_demo},"Confirmed limitation")+COUNTIFS({rng_area},A{row},{rng_demo},"Security finding")')
        s.cell(row=row, column=5, value=f'=COUNTIFS({rng_area},A{row},{rng_demo},"Not run")')
        s.cell(row=row, column=6, value=f'=COUNTIFS({rng_area},A{row},{rng_mine},"Pass")')
        s.cell(row=row, column=7, value=f'=COUNTIFS({rng_area},A{row},{rng_mine},"Fail")')
        row += 1
    s.cell(row=row, column=1, value="Total")
    for col in "BCDEFG":
        s[f"{col}{row}"] = f"=SUM({col}5:{col}{row - 1})"
    for rr in range(5, row + 1):
        for col in range(1, 8):
            cell = s.cell(row=rr, column=col)
            cell.font = _font(bold=rr == row)
            cell.border = BORDER
            if rr == row:
                cell.fill = SOFT_FILL
    total = row
    row += 2
    _table(s, row, ["Your progress", "Value"], [], [30, 15])
    items = [("Scenarios with a result", f'=COUNTA({rng_mine})'),
             ("Passed", f'=COUNTIF({rng_mine},"Pass")'),
             ("Failed", f'=COUNTIF({rng_mine},"Fail")'),
             ("P1 scenarios", f'=COUNTIF({rng_prio},"P1")'),
             ("P1 scenarios passed", f'=COUNTIFS({rng_prio},"P1",{rng_mine},"Pass")'),
             ("P1 pass rate", f'=IFERROR(B{row + 5}/B{row + 4},0)')]
    for i, (label, formula) in enumerate(items, row + 1):
        s.cell(row=i, column=1, value=label).font = _font()
        cell = s.cell(row=i, column=2, value=formula)
        cell.font = _font()
        if label == "P1 pass rate":
            cell.number_format = "0.0%"
        for col in (1, 2):
            s.cell(row=i, column=col).border = BORDER
    s.cell(row=row + len(items) + 2, column=1,
           value=f"Demo tenant totals: see row {total}. Use cases and the persona matrix are on the other sheets.").font = _font(color="5C5C5C")
    s.row_dimensions[4].height = 30
    _fit_rows(tests, 2, last, header_row=1)
    _fit_rows(usecases, 2, last_uc, header_row=1)
    _print_setup(tests, "1:1")
    _print_setup(usecases, "1:1")
    _print_setup(summary)
    for ws in (summary, tests, usecases, setup):
        ws.sheet_view.showGridLines = ws is not summary
    wb.save(path)


def catalog():
    by_area = []
    for area in c.AREAS:
        items = [t for t in c.TESTS if t["area"] == area]
        by_area.append({"area": area, "total": len(items), **{k: sum(t["status"] == k for t in items) for k in c.STATUS}})
    return {"tested_on": c.TESTED_ON, "package_version": c.PACKAGE_VERSION, "stats": c.S, "counts": c.counts(),
            "by_area": by_area, "core_set": c.CORE_SET,
            "use_cases": [{k: u[k] for k in ("id", "title", "status", "scenario", "how", "evidence", "tests")} for u in c.USE_CASES],
            "tests": c.TESTS}


def main():
    c.check()
    with open(os.path.join(DOCS, TUTORIAL_MD), "w", encoding="utf-8", newline="\n") as f:
        f.write(tutorial_md())
    with open(os.path.join(DOCS, SCENARIOS_MD), "w", encoding="utf-8", newline="\n") as f:
        f.write(scenarios_md())
    with open(os.path.join(DOCS, CATALOG), "w", encoding="utf-8", newline="\n") as f:
        json.dump(catalog(), f, indent=1, ensure_ascii=False)
    workbook(os.path.join(DOCS, XLSX))
    print("wrote", TUTORIAL_MD, SCENARIOS_MD, XLSX, CATALOG)


if __name__ == "__main__":
    main()
